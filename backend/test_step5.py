"""
STEP 5 Verification & End-to-End Test Script
Starts real Uvicorn server, creates real PDF, DOCX, XLSX, CSV, and PNG test files,
uploads & processes them via real HTTP requests, and verifies PostgreSQL persistence & file integrity.
"""

import os
import sys
import time
import subprocess
import json
import fitz
import docx
import pandas as pd
from PIL import Image, ImageDraw
import httpx

BASE_URL = "http://127.0.0.1:8000"
TEST_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_files")
os.makedirs(TEST_DIR, exist_ok=True)


def create_test_pdf():
    pdf_path = os.path.join(TEST_DIR, "sample_coal_report.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        "CMPDI GEOLOGICAL EXPLORATION REPORT 2026\n\n"
        "Block Name: North Karanpura Block C\n"
        "Coal Seam: Seam VI (Top)\n"
        "Thickness: 5.2 meters\n"
        "Ash Content: 17.5%\n"
        "Gross Calorific Value (GCV): 5400 kcal/kg\n"
        "Proved Reserves: 14.8 Million Tonnes\n"
    )
    doc.save(pdf_path)
    doc.close()
    return pdf_path


def create_test_docx():
    docx_path = os.path.join(TEST_DIR, "sample_core_log.docx")
    doc = docx.Document()
    doc.add_heading("CMPDI Core Drilling Log - Borehole NC-12", level=1)
    doc.add_paragraph("Location: Rajrappa Area, Central Coalfields Limited.")
    doc.add_paragraph("Drilling Date: 2026-08-15. Target Seam: Argada Seam.")
    
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "From_m"
    hdr_cells[1].text = "To_m"
    hdr_cells[2].text = "Lithology"
    hdr_cells[3].text = "Seam_Code"
    
    row_data = [
        ("0.0", "15.2", "Overburden Sandstone", "OB"),
        ("15.2", "22.8", "Bituminous Coal", "ARG-1"),
        ("22.8", "35.0", "Shale & Carbonaceous Shale", "SH"),
    ]
    for from_m, to_m, lith, seam in row_data:
        row_cells = table.add_row().cells
        row_cells[0].text = from_m
        row_cells[1].text = to_m
        row_cells[2].text = lith
        row_cells[3].text = seam
        
    doc.save(docx_path)
    return docx_path


def create_test_xlsx():
    xlsx_path = os.path.join(TEST_DIR, "sample_production.xlsx")
    df_prod = pd.DataFrame({
        "Mine_Name": ["Amrapali OCP", "Magadh OCP", "Ashoka OCP"],
        "Target_MT": [12.5, 10.0, 8.5],
        "Actual_MT": [13.1, 9.8, 8.7],
        "Achievement_Pct": [104.8, 98.0, 102.35]
    })
    df_quality = pd.DataFrame({
        "Grade": ["G7", "G8", "G10"],
        "GCV_Band": ["5201-5500", "4901-5200", "4301-4600"],
        "Dispatch_MT": [5.2, 8.1, 4.0]
    })
    
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df_prod.to_excel(writer, sheet_name="Production_Summary", index=False)
        df_quality.to_excel(writer, sheet_name="Quality_Grade", index=False)
        
    return xlsx_path


def create_test_csv():
    csv_path = os.path.join(TEST_DIR, "sample_boreholes.csv")
    df = pd.DataFrame({
        "Borehole_ID": ["BH-01", "BH-02", "BH-03"],
        "Collar_RL": [245.5, 250.1, 242.8],
        "Total_Depth_m": [180.0, 210.5, 165.0],
        "Seam_Count": [4, 5, 3]
    })
    df.to_csv(csv_path, index=False)
    return csv_path


def create_test_image():
    img_path = os.path.join(TEST_DIR, "sample_map.png")
    img = Image.new("RGB", (400, 200), color=(240, 240, 240))
    d = ImageDraw.Draw(img)
    d.text((20, 80), "CMPDI MINING SCHEME BLOCK B", fill=(0, 0, 0))
    img.save(img_path)
    return img_path


def run_tests():
    print("=== STARTING REAL UVICORN BACKEND PROCESS FOR STEP 5 TEST ===", flush=True)
    
    python_bin = sys.executable
    server_proc = subprocess.Popen(
        [python_bin, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", "8000"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL
    )
    
    # Wait for server startup
    time.sleep(3)
    
    try:
        with httpx.Client(base_url=BASE_URL, timeout=120.0) as client:
            # 1. Health checks
            r = client.get("/api/health")
            assert r.status_code == 200, f"Health failed: {r.text}"
            print("[PASS] GET /api/health:", r.json())
            
            r_db = client.get("/api/health/db")
            assert r_db.status_code == 200, f"DB Health failed: {r_db.text}"
            print("[PASS] GET /api/health/db:", r_db.json())
            
            # 2. Prepare real files
            pdf_path = create_test_pdf()
            docx_path = create_test_docx()
            xlsx_path = create_test_xlsx()
            csv_path = create_test_csv()
            img_path = create_test_image()
            
            test_files = [
                ("PDF Report", "pdf", pdf_path, "application/pdf"),
                ("DOCX Core Log", "docx", docx_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                ("XLSX Production", "xlsx", xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("CSV Boreholes", "csv", csv_path, "text/csv"),
                ("PNG Mining Map", "png", img_path, "image/png"),
            ]
            
            processed_results = []
            
            for name, ftype, fpath, ctype in test_files:
                print(f"\n--- Testing file type: {ftype.upper()} ({os.path.basename(fpath)}) ---")
                
                orig_stat = os.stat(fpath)
                
                with open(fpath, "rb") as f:
                    resp = client.post(
                        "/api/documents/upload",
                        files={"file": (os.path.basename(fpath), f, ctype)},
                        data={
                            "name": name,
                            "type": ftype,
                            "source": "CMPDI Test Suite",
                            "category": "Geological Data",
                            "date": "2026-09-16",
                            "description": f"Automated test upload for {ftype}",
                        }
                    )
                    
                assert resp.status_code == 201, f"Upload failed ({resp.status_code}): {resp.text}"
                data = resp.json()
                doc_id = data["document"]["id"]
                status = data["document"]["processing_status"]
                print(f"[PASS] Upload & Auto-processing response for doc ID {doc_id}: status='{status}'")
                
                # Fetch content via GET /api/documents/{id}/content
                content_resp = client.get(f"/api/documents/{doc_id}/content")
                assert content_resp.status_code == 200, f"Get content failed: {content_resp.text}"
                content_data = content_resp.json()
                
                assert content_data["processing_status"] == "completed", f"Processing status not completed: {content_data.get('error_message')}"
                assert len(content_data["extracted_text"]) > 0, "Extracted text is empty!"
                assert len(content_data["chunks"]) > 0, "No chunks generated!"
                
                print(f"  - Extracted Text Length: {len(content_data['extracted_text'])} chars")
                print(f"  - Chunks Stored in DB: {len(content_data['chunks'])}")
                print(f"  - Extracted Text Snippet:\n    {content_data['extracted_text'][:150]!r}")
                
                # Verify re-processing endpoint POST /api/documents/{id}/process
                proc_resp = client.post(f"/api/documents/{doc_id}/process")
                assert proc_resp.status_code == 200, f"Re-processing failed: {proc_resp.text}"
                print(f"[PASS] Explicit POST /api/documents/{doc_id}/process endpoint call succeeded.")
                
                # Verify original file on disk was NOT modified or deleted
                post_stat = os.stat(fpath)
                assert orig_stat.st_size == post_stat.st_size, "Original file size changed!"
                assert orig_stat.st_mtime == post_stat.st_mtime, "Original file was modified!"
                print(f"[PASS] Original file '{os.path.basename(fpath)}' remains strictly intact (size={post_stat.st_size} bytes).")
                
                processed_results.append({
                    "id": doc_id,
                    "filename": os.path.basename(fpath),
                    "type": ftype,
                    "status": content_data["processing_status"],
                    "page_count": content_data["page_count"],
                    "chunks": len(content_data["chunks"])
                })
                
            print("\n=== ALL 5 FILE TYPES PROCESSED & PERSISTED SUCCESSFULLY IN POSTGRESQL ===")
            print(json.dumps(processed_results, indent=2))
            
    finally:
        server_proc.terminate()
        server_proc.wait()
        print("\nUvicorn backend test process terminated cleanly.")


if __name__ == "__main__":
    run_tests()
